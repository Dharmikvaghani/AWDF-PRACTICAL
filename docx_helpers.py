import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    val: "single", sz: "4", color: "000000"
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', '000000'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)
        
        # Header
        header = s.header
        header.is_linked_to_previous = False
        p_hdr = header.paragraphs[0]
        p_hdr.text = ""
        tbl_hdr = header.add_table(1, 2, Inches(7.2))
        tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_hdr.rows[0].cells
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run("ITUE301: Advanced Web Development Frameworks")
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run("24CS104")
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.bold = True
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        # Footer
        footer = s.footer
        footer.is_linked_to_previous = False
        p_ftr = footer.paragraphs[0]
        p_ftr.text = ""
        
        # Add a decorative colored rule
        p_rule = footer.add_paragraph()
        p_rule.paragraph_format.space_before = Pt(0)
        p_rule.paragraph_format.space_after = Pt(2)
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="622423"/></w:pBdr>')
        p_rule._p.get_or_add_pPr().append(pBdr)
        
        tbl_ftr = footer.add_table(1, 2, Inches(7.2))
        tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_ftr.rows[0].cells
        
        fp0 = fc0.paragraphs[0]
        fr0 = fp0.add_run("Department of Computer Science and Engineering, CSPIT")
        fr0.font.name = "Calibri"
        fr0.font.size = Pt(9.5)
        fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        fp1 = fc1.paragraphs[0]
        fr1 = fp1.add_run("Page ")
        # Page field
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp1._p.append(fldSimple)
        fr1.font.name = "Calibri"
        fr1.font.size = Pt(9.5)
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def create_meta_table(doc, practical_num):
    table = doc.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="000000", sz="4")
    
    # Col widths: total 7.2 in
    col_widths = [Inches(1.2), Inches(1.3), Inches(1.3), Inches(0.7), Inches(1.5), Inches(1.2)]
    
    # Row 0: Student ID : 24CS104 | Student Name: Dharmik Vaghani
    row0 = table.rows[0]
    # merge col 1 and 2 for Student ID, or merge cols for Name
    # Student Details: [Student ID :] [24CS104] [Student Name:] [Dharmik Vaghani]
    # Let's merge cell 0 & 1? Let's check:
    # Cell 0: Student ID :
    # Cell 1: 24CS104
    # Cell 2: Student Name:
    # Cell 3-5 merged: Dharmik Vaghani
    c_id_label = row0.cells[0]
    c_id_val = row0.cells[1]
    c_name_label = row0.cells[2]
    c_name_val = row0.cells[3]
    c_name_val.merge(row0.cells[4]).merge(row0.cells[5])
    
    p = c_id_label.paragraphs[0]
    r = p.add_run("Student ID :")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    
    p = c_id_val.paragraphs[0]
    r = p.add_run("24CS104")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    
    p = c_name_label.paragraphs[0]
    r = p.add_run("Student Name:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    
    p = c_name_val.paragraphs[0]
    r = p.add_run("Dharmik Vaghani")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    
    # Row 1: Subject code & Name : | ITUE301: Advanced Web Development Frameworks | Practical : | {practical_num} | Academic Year : | 2026-27
    row1 = table.rows[1]
    row1.cells[0].paragraphs[0].add_run("Subject code &\nName :").font.size = Pt(9.5)
    row1.cells[1].paragraphs[0].add_run("ITUE301: Advanced Web\nDevelopment\nFrameworks").font.size = Pt(9)
    row1.cells[2].paragraphs[0].add_run("Practical :").font.size = Pt(9.5)
    row1.cells[3].paragraphs[0].add_run(str(practical_num)).font.size = Pt(10)
    row1.cells[4].paragraphs[0].add_run("Academic Year :").font.size = Pt(9.5)
    row1.cells[5].paragraphs[0].add_run("2026-27").font.size = Pt(10)
    
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for cp in cell.paragraphs:
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.line_spacing = 1.05
    return table

print('Helper definitions loaded!')
