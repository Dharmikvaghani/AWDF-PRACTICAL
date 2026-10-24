import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def setup_page_layout(section):
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    
    # Header
    header = section.header
    header.is_linked_to_previous = False
    p_hdr = header.paragraphs[0]
    p_hdr.text = ""
    tbl_hdr = header.add_table(1, 2, Inches(7.2))
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = tbl_hdr.rows[0].cells
    
    p0 = c0.paragraphs[0]
    r0 = p0.add_run("ITUE301: Advanced Web Development Frameworks")
    r0.font.name = "Calibri"
    r0.font.size = Pt(10)
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    p1 = c1.paragraphs[0]
    r1 = p1.add_run("24CS104")
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    r1.font.bold = True
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    p_ftr = footer.paragraphs[0]
    p_ftr.text = ""
    
    # Maroon bottom rule
    p_rule = footer.add_paragraph()
    p_rule.paragraph_format.space_before = Pt(0)
    p_rule.paragraph_format.space_after = Pt(2)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="622423"/></w:pBdr>')
    p_rule._p.get_or_add_pPr().append(pBdr)
    
    tbl_ftr = footer.add_table(1, 2, Inches(7.2))
    tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
    fc0, fc1 = tbl_ftr.rows[0].cells
    
    fp0 = fc0.paragraphs[0]
    fr0 = fp0.add_run("Department of Computer Science and Engineering, CSPIT")
    fr0.font.name = "Calibri"
    fr0.font.size = Pt(9.5)
    fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    fp1 = fc1.paragraphs[0]
    fr1 = fp1.add_run("Page ")
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp1._p.append(fldSimple)
    fr1.font.name = "Calibri"
    fr1.font.size = Pt(9.5)
    fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def create_doc():
    doc = docx.Document()
    setup_page_layout(doc.sections[0])
    
    with open('all_practicals_data.json', encoding='utf-8') as f:
        practicals = json.load(f)
        
    for idx, item in enumerate(practicals):
        p_num = item['practical_num']
        
        # ================= PAGE 1 =================
        # 1. Meta Table
        table_meta = doc.add_table(rows=2, cols=6)
        table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_meta.autofit = False
        set_table_borders(table_meta, color="000000", sz="4")
        
        # Row 0
        r0 = table_meta.rows[0]
        r0.cells[3].merge(r0.cells[4]).merge(r0.cells[5])
        
        p = r0.cells[0].paragraphs[0]
        p.add_run("Student ID :").font.size = Pt(10)
        p = r0.cells[1].paragraphs[0]
        r = p.add_run("24CS104")
        r.bold = True
        r.font.size = Pt(10)
        
        p = r0.cells[2].paragraphs[0]
        p.add_run("Student Name:").font.size = Pt(10)
        p = r0.cells[3].paragraphs[0]
        r = p.add_run("Dharmik Vaghani")
        r.bold = True
        r.font.size = Pt(10)
        
        # Row 1
        r1 = table_meta.rows[1]
        p = r1.cells[0].paragraphs[0]
        p.add_run("Subject code &\nName :").font.size = Pt(9.5)
        p = r1.cells[1].paragraphs[0]
        p.add_run("ITUE301: Advanced Web\nDevelopment\nFrameworks").font.size = Pt(9)
        p = r1.cells[2].paragraphs[0]
        p.add_run("Practical :").font.size = Pt(9.5)
        p = r1.cells[3].paragraphs[0]
        p.add_run(str(p_num)).font.size = Pt(10)
        p = r1.cells[4].paragraphs[0]
        p.add_run("Academic Year :").font.size = Pt(9.5)
        p = r1.cells[5].paragraphs[0]
        p.add_run("2026-27").font.size = Pt(10)
        
        for row in table_meta.rows:
            for cell in row.cells:
                set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for cp in cell.paragraphs:
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = 1.05
                    for run in cp.runs:
                        run.font.name = "Calibri"

        # Space between meta table and content table
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_before = Pt(4)
        sp_p.paragraph_format.space_after = Pt(4)
        sp_p.paragraph_format.line_spacing = 1.0
        
        # 2. Main Content Box (Page 1)
        # Rows: Task Header, Aim Header, Aim Text, Code Header, Code Content
        content_table = doc.add_table(rows=5, cols=1)
        content_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        content_table.autofit = False
        set_table_borders(content_table, color="000000", sz="4")
        
        # Row 0: Task - X
        c = content_table.rows[0].cells[0]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Task - {p_num}")
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        set_cell_margins(c, top=40, bottom=40, left=100, right=100)
        
        # Row 1: Aim
        c = content_table.rows[1].cells[0]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Aim")
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        set_cell_margins(c, top=40, bottom=40, left=100, right=100)
        
        # Row 2: Aim text
        c = content_table.rows[2].cells[0]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(item['aim'])
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        set_cell_margins(c, top=50, bottom=50, left=100, right=100)
        
        # Row 3: Code
        c = content_table.rows[3].cells[0]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Code")
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        set_cell_margins(c, top=40, bottom=40, left=100, right=100)
        
        # Row 4: Code Content & Highlights
        c = content_table.rows[4].cells[0]
        set_cell_margins(c, top=60, bottom=60, left=100, right=100)
        p_first = c.paragraphs[0]
        p_first.paragraph_format.space_before = Pt(0)
        p_first.paragraph_format.space_after = Pt(2)
        r_gh_label = p_first.add_run("GitHub Link / Repository Link: ")
        r_gh_label.font.name = "Calibri"
        r_gh_label.font.size = Pt(9.5)
        r_gh_label.bold = True
        
        r_gh_val = p_first.add_run(item['github_link'])
        r_gh_val.font.name = "Calibri"
        r_gh_val.font.size = Pt(9.5)
        r_gh_val.font.color.rgb = RGBColor(37, 99, 235)
        r_gh_val.font.underline = True
        
        # Highlights Header
        p_hl_hdr = c.add_paragraph()
        p_hl_hdr.paragraph_format.space_before = Pt(4)
        p_hl_hdr.paragraph_format.space_after = Pt(2)
        r = p_hl_hdr.add_run(item['highlights_title'])
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        
        # Highlights items
        for hl in item['highlights']:
            p_hl = c.add_paragraph()
            p_hl.paragraph_format.space_before = Pt(0)
            p_hl.paragraph_format.space_after = Pt(1)
            p_hl.paragraph_format.line_spacing = 1.05
            r = p_hl.add_run(hl)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            
        # Code snippets
        if item['code']:
            p_code_gap = c.add_paragraph()
            p_code_gap.paragraph_format.space_before = Pt(3)
            p_code_gap.paragraph_format.space_after = Pt(0)
            
            code_lines = item['code'].split('\n')
            for cl in code_lines:
                p_code = c.add_paragraph()
                p_code.paragraph_format.space_before = Pt(0)
                p_code.paragraph_format.space_after = Pt(0)
                p_code.paragraph_format.line_spacing = 1.05
                r = p_code.add_run(cl)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                # color comments gray
                if cl.strip().startswith('//') or cl.strip().startswith('#'):
                    r.font.color.rgb = RGBColor(100, 116, 139)
                else:
                    r.font.color.rgb = RGBColor(15, 23, 42)
                    
        # ================= PAGE 2 =================
        doc.add_page_break()
        
        # Page 2 content table
        p2_rows = 2
        if item.get('page2_code'):
            p2_rows = 3
            
        tbl_p2 = doc.add_table(rows=p2_rows, cols=1)
        tbl_p2.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_p2.autofit = False
        set_table_borders(tbl_p2, color="000000", sz="4")
        
        curr_row = 0
        if item.get('page2_code'):
            c_p2c = tbl_p2.rows[curr_row].cells[0]
            set_cell_margins(c_p2c, top=40, bottom=40, left=100, right=100)
            for cl in item['page2_code'].split('\n'):
                p_c = c_p2c.add_paragraph() if c_p2c.paragraphs[0].text else c_p2c.paragraphs[0]
                p_c.paragraph_format.space_before = Pt(0)
                p_c.paragraph_format.space_after = Pt(0)
                p_c.paragraph_format.line_spacing = 1.05
                r = p_c.add_run(cl)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
            curr_row += 1
            
        # Output Header row
        c_out_hdr = tbl_p2.rows[curr_row].cells[0]
        set_cell_margins(c_out_hdr, top=40, bottom=40, left=100, right=100)
        p = c_out_hdr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Output")
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        curr_row += 1
        
        # Image row
        c_img = tbl_p2.rows[curr_row].cells[0]
        set_cell_margins(c_img, top=100, bottom=100, left=80, right=80)
        p_img = c_img.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(6)
        
        fig_path = item['figure_img']
        if os.path.exists(fig_path):
            p_img.add_run().add_picture(fig_path, width=Inches(6.2))
            
        # Caption outside/below table
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(item['figure_caption'])
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        
        # If not the last practical, page break to start next practical
        if idx < len(practicals) - 1:
            doc.add_page_break()

    out_name = "AWDF_Practical_24CS104_Dharmik_Vaghani.docx"
    doc.save(out_name)
    print(f"Successfully generated {out_name}!")

if __name__ == "__main__":
    create_doc()
